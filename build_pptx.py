#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
参数单生成器 V2 — 后端最小闭环
读数据 JSON -> 套模板 (template_map.<lang>.json) -> 出 pptx

用法:
  python3 build_pptx.py --data data/D17.json --lang en --out output/
"""
import argparse
import copy
import io
import json
import os
import re
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))


def load_map(lang):
    p = os.path.join(HERE, "templates", "template_map.%s.json" % lang)
    with io.open(p, encoding="utf-8") as f:
        return json.load(f)


_GLOSSARY = None


def load_glossary():
    global _GLOSSARY
    if _GLOSSARY is None:
        p = os.path.join(HERE, "glossary.json")
        if os.path.exists(p):
            with io.open(p, encoding="utf-8") as f:
                raw = json.load(f)
            # rows 格式: [中文, 英文, 适用]，忽略以 __ 开头的分类标题行
            zh2en, en2zh, prefix = {}, {}, {}
            for row in raw.get("rows", []):
                if len(row) < 2:
                    continue
                zh, en = row[0], row[1]
                if not zh or zh.startswith("__"):
                    continue
                zh2en[zh] = en
                if en and en not in en2zh:
                    en2zh[en] = zh
                if zh.endswith(" "):
                    prefix[zh] = en
            _GLOSSARY = {"zh2en": zh2en, "en2zh": en2zh, "prefix": prefix}
        else:
            _GLOSSARY = {"zh2en": {}, "en2zh": {}, "prefix": {}}
    return _GLOSSARY


def translate(val, lang):
    """按语言翻译单元格值。

    - zh: 若值是英文术语（命中 en2zh）则转回中文，否则原样
    - en: 若值是中文术语（命中 zh2en）则转英文，否则原样
    - 整值精确命中优先；否则做前缀替换（如 '等级 H/H' -> 'Class H/H'）
    - 数值/含单位可变值不命中则不动
    """
    if val is None:
        return val
    g = load_glossary()
    s = str(val)
    if lang == "en":
        if s in g["zh2en"]:
            return g["zh2en"][s]
        for k, v in g["prefix"].items():
            if s.startswith(k):
                return v + s[len(k):]
    elif lang == "zh":
        if s in g["en2zh"]:
            return g["en2zh"][s]
    return val


def _clean_value(v):
    """轻度清洗写入值：保留行首 `◆`/缩进等排版符号，只把制表符与不换行空格归一。

    注意：绝不能 strip 行首 —— features 行的 `◆ `、缩进都是模板设计要求，
    去掉会让文字挤成一团、断行错乱。只做：\xa0->空格、去行尾空白、去空行。
    """
    if v is None:
        return None
    s = str(v).replace("\u00a0", " ")
    lines = []
    for x in s.split("\n"):
        x = x.rstrip()          # 只去行尾空白，保留行首
        lines.append(x)
    # 去首尾空行，保留中间空行
    while lines and lines[0] == "":
        lines.pop(0)
    while lines and lines[-1] == "":
        lines.pop()
    return "\n".join(lines)


def _text_width_in(text, pt, font_path=None):
    """估算文本宽度（英寸）。用 Pillow + Bai Jamjuree 真实字形宽度。"""
    try:
        from PIL import ImageFont
        px = max(1, int(round(pt * 96 / 72)))
        path = font_path or "/usr/share/fonts/baijamjuree/BaiJamjuree-Regular.ttf"
        if not os.path.exists(path):
            path = "/usr/share/fonts/baijamjuree/BaiJamjuree-Medium.ttf"
        f = ImageFont.truetype(path, px)
        return f.getlength(text) / 96.0
    except Exception:
        # 兜底：按 8pt 平均字宽 0.055in 粗估
        return len(text) * 0.055 * (pt / 8.0)


# 中文版用什么字体量宽：Bai Jamjuree 无中文字形，LibreOffice 走 fontconfig 回落
# Noto Sans CJK SC（服务器 fonts.conf 已把 黑体/SimHei 指过去）。
_CJK_FONT_CANDIDATES = [
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc",
    "/System/Library/Fonts/PingFang.ttc",
]


def _has_cjk(text):
    return any("\u2e80" <= ch <= "\u9fff" or "\uff00" <= ch <= "\uffef" for ch in str(text))


def _cjk_font_path():
    for p in _CJK_FONT_CANDIDATES:
        if os.path.exists(p):
            return p
    return None


def _shrink_to_fit(cell, text, min_pt=5.5):
    """确保单元格内文本不折行：按可用宽度自动降字号。

    可用宽度 = 列宽 - 左右内边距（模板多为 0.08in）。
    只调字号（视觉几乎看不出差异），不改列宽、不动版式。
    含中日文字符时改用 CJK 字体量宽（Bai Jamjuree 没有中文字形，
    用它量中文会严重低估宽度，导致照旧折行）。
    """
    font_path = _cjk_font_path() if _has_cjk(text) else None
    try:
        from pptx.util import Pt as _Pt
        col_w = None
        # 通过父 table 反查列宽
        tc = cell._tc
        tr = tc.getparent()
        tbl = tr.getparent()
        cols = tbl.findall(qn("a:tblGrid") + "/" + qn("a:gridCol"))
        # 定位当前列序
        idx = list(tr).index(tc)
        if idx < len(cols):
            col_w = int(cols[idx].get("w", 0)) / 914400.0
        if not col_w:
            return
        marL = (cell.margin_left or 0) / 914400.0
        marR = (cell.margin_right or 0) / 914400.0
        avail = col_w - marL - marR
        # 取当前字号
        cur = None
        for para in cell.text_frame.paragraphs:
            for r in para.runs:
                if r.font.size:
                    cur = r.font.size.pt
                    break
            if cur:
                break
        if cur is None:
            cur = 8.0
        # Pillow 估宽偏乐观（LibreOffice 实际渲染更宽），留 12% 安全余量
        avail = avail * 0.88
        pt = cur
        while pt > min_pt and _text_width_in(text, pt, font_path) > avail:
            pt -= 0.5
        if pt < cur:
            for para in cell.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = _Pt(pt)
    except Exception as e:
        print("  ! 缩字号失败:", e, file=sys.stderr)


# 需要「一格一行」的表格：第1页基础参数表、第2页全部参数表（含表头与标签列）
NOWRAP_PURPOSES = {
    "power_rating", "performance_data", "dimensions",
    "engine_basic_data", "exhaust_system", "fuel_system", "cooling_system",
    "electrical_system", "inlet_air", "alternator_basic_data", "alternator_structure",
}


def _cell_width_in(cell):
    """列宽（英寸），已扣左右内边距。"""
    tc = cell._tc
    tr = tc.getparent()
    tbl = tr.getparent()
    cols = tbl.findall(qn("a:tblGrid") + "/" + qn("a:gridCol"))
    idx = list(tr).index(tc)
    if idx >= len(cols):
        return None
    w = int(cols[idx].get("w", 0)) / 914400.0
    return w - (cell.margin_left or 0) / 914400.0 - (cell.margin_right or 0) / 914400.0


_RENDER_PROBE_ENABLED = [True]


def _fit_cell_one_line(cell, min_pt):
    """把单元格文字压成一行：降字号优先，降到底还放不下才按单词折行。"""
    try:
        from pptx.util import Pt as _Pt
        runs = [r for p in cell.text_frame.paragraphs for r in p.runs]
        if not runs:
            return None, None
        text = "".join(r.text or "" for r in runs)
        if not text.strip():
            return 8.0, None
        avail = _cell_width_in(cell)
        if not avail:
            return None, None
        avail *= 0.88          # Pillow 估宽偏乐观，留 12% 余量
        cur = next((r.font.size.pt for r in runs if r.font.size), 8.0)
        fp = _cjk_font_path() if _has_cjk(text) else None

        # ⚠️ 跨平台误差：同一 TTF 在 macOS / Linux 上 getlength 结果能差 10%+
        # （Bai Jamjuree「Prime & Standby」8pt：Mac 0.88in / Linux 0.97in），
        # 纯靠估宽会算出「放得下」但 PDF 照样折行。所以在服务器上直接写字量：
        # 把文字临时写到单元格里，让 LibreOffice 摊出真实单行宽度，
        # 超出可用宽度就降半磅重试（最多降 8 次），然后恢复原文。
        def _fit_by_render(text, start_pt):
            try:
                import subprocess, tempfile, shutil
                sofa = shutil.which("soffice") or shutil.which("libreoffice")
                if not sofa or not _RENDER_PROBE_ENABLED[0]:
                    return None
                # 实测渲染宽度时要带上 12% 安全余量，所以探针框也取同一宽度
                usable_emu = int(round(avail * 914400))
                # 用一个空白 slide：单元格自动换行时，摊出的「单行宽度」就等于可用宽度
                tmpd = tempfile.mkdtemp()
                try:
                    probe = Presentation()
                    blank = probe.slide_layouts[6]
                    sl = probe.slides.add_slide(blank)
                    tbl = sl.shapes.add_table(1, 1, 0, 0, usable_emu, 914400).table
                    tcell = tbl.cell(0, 0)
                    tcell.margin_left = tcell.margin_right = 0
                    pt = float(start_pt)
                    for _ in range(9):
                        tcell.text = text
                        for para in tcell.text_frame.paragraphs:
                            for r in para.runs:
                                r.font.size = _Pt(pt)
                                r.font.name = "Bai Jamjuree"
                        pp = os.path.join(tmpd, "p.pptx")
                        probe.save(pp)
                        subprocess.run([sofa, "--headless", "--convert-to", "pdf",
                                        "--outdir", tmpd, pp],
                                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                       timeout=60)
                        pdff = os.path.join(tmpd, "p.pdf")
                        if not os.path.exists(pdff):
                            return None
                        import pdfplumber
                        with pdfplumber.open(pdff) as pf:
                            words = pf.pages[0].extract_words()
                        if not words:
                            return None
                        used = max(w["x1"] for w in words) - min(w["x0"] for w in words)
                        # 摊出宽度比可用宽度还小 → 没折行，放得下
                        if used <= (usable_emu / 914400.0) * 0.985:
                            return pt
                        pt -= 0.5
                    return None
                finally:
                    shutil.rmtree(tmpd, ignore_errors=True)
            except Exception as e:
                print("  ! 渲染量宽失败:", e, file=sys.stderr)
                return None

        pt = cur
        while pt > min_pt and _text_width_in(text, pt, fp) > avail:
            pt -= 0.5
        # 2026-10-20 关掉「转 PDF 实测复核」：模板里同一格 LibreOffice 有 0.74in 的
        # 隐形最小宽度，单格探针测出来的数不可信，还会让每次导出多花十几秒。
        # 现在靠「缩字号 + 列宽让位」解决（见 _FUEL_COL_FIX）。
        if False and _RENDER_PROBE_ENABLED[0] and not _has_cjk(text):
            real = _fit_by_render(text, pt)
            _RENDER_PROBE_ENABLED[0] = False    # 探一次即够，后续表用同一套量宽结果
            if real is not None:
                pt = min(pt, real)
        if pt < cur:
            for r in runs:
                r.font.size = _Pt(pt)
        # 到下限仍放不下 → 用单词软换行（英文在空格处折，CJK 用 U+200B 断点）
        if pt <= min_pt and _text_width_in(text, pt, fp) > avail:
            if " " in text:
                words, lines, cur_line = text.split(" "), [], ""
                for wd in words:
                    trial = (cur_line + " " + wd) if cur_line else wd
                    if _text_width_in(trial, pt, fp) <= avail or not cur_line:
                        cur_line = trial
                    else:
                        lines.append(cur_line)
                        cur_line = wd
                if cur_line:
                    lines.append(cur_line)
            else:
                lines, cur_line = [], ""
                for ch in text:
                    if _text_width_in(cur_line + ch, pt, fp) <= avail or not cur_line:
                        cur_line += ch
                    else:
                        lines.append(cur_line)
                        cur_line = ch
                if cur_line:
                    lines.append(cur_line)
            if len(lines) > 1:
                set_multiline(cell, lines)
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = _Pt(pt)
                return pt, lines
        return pt, None
    except Exception as e:
        print("  ! 单行压缩失败:", e, file=sys.stderr)
        return None, None


def _nowrap_table(table):
    """整个表格的单元格一律不折行（表头/标签/数据格）。

    ⚠️ 表头和标签列是模板原文、**不能动**（用户铁律：模板是权威），
    只有<b>数据格</b>才允许缩字号/折词：即「模板里原文是 12345 占位」的格子。
    判定：去掉空白后的原文全部由数字组成 → 视为占位数据格。
    """
    for r in range(len(table.rows)):
        for c in range(len(table.columns)):
            cell = table.cell(r, c)
            if not cell.text.strip():
                continue
            # 表头/标签是模板原文 → 一律不碰（模板是权威）
            if not re.fullmatch(r"\s*\d+\s*", cell.text):
                continue
            _fit_cell_one_line(cell, min_pt=5.0)


def set_cell(cell, text, size_pt=None, nowrap=False, min_pt=5.5):
    """写单元格，保留原格式（清空多余 run/段落）

    size_pt: 可选，覆盖字号（如尺寸列 13 字符在 1.1in 宽 8pt 下会折行，降到 7pt）。
    nowrap:  可选，禁止单元格内自动换行（长数字如 3200*1365*1820、
             2000 kVA 不要被拆成两行）。
    min_pt:  nowrap 生效时的最小字号。模板表头/标签文字固定、不能再缩，
             所以表头用 6.5；数据格允许降到 5.0。
    """
    if text is None:
        return
    text = _clean_value(text)
    set_multiline(cell, [str(text)])
    if size_pt:
        for para in cell.text_frame.paragraphs:
            for r in para.runs:
                r.font.size = Pt(size_pt)
    if nowrap:
        # ⚠️ LibreOffice 转 PDF 时对表格单元格**忽略** bodyPr wrap="none"，
        # 所以真正防折行只能缩字号：按可用宽度自动往下试。
        _shrink_to_fit(cell, str(text), min_pt=min_pt)


def _para_prefix(p_el):
    """读出模板段落首个 run 的『符号前缀』（如 '◆ '）。

    模板 features 行首个 run 常是 '◆ '（符号单独成 run）；重建段落时必须
    原样保留，否则 ◆ 会被用户数据覆盖掉，导致文字挤成一团、断行错乱。
    判定：首个 run 文本去掉符号/空白后为空 → 视为纯前缀，返回 (前缀, 是否独立)。
    """
    ts = p_el.findall(qn("a:r"))
    if not ts:
        return None
    t = ts[0].find(qn("a:t"))
    txt = t.text if (t is not None and t.text) else ""
    core = txt.replace("◆", "").replace("●", "").replace("•", "").strip()
    if core == "":           # 纯符号/空白前缀
        return txt
    return None


def set_multiline(cell, lines, keep_prefix=False, keep_shape=False):
    """多行写入：清空后按模板首个段落样式重建所有行。

    cell 可为表格单元格，也可为 shape.text_frame（keep_shape=True 时）。
    必须先在副本上改文本、再整体替换原始 <a:p>，否则按 _p 查找会重复命中。
    空行也必须保留为独立段落（否则会多出/丢掉行，如油耗 4 行夹着空值）。
    keep_prefix=True 时，为每行自动补上模板的符号前缀（如 '◆ '），
    用于 features 这类「每行一个 ◆」的单元格。
    """
    if not lines:
        return
    lines = ["" if x is None else str(x) for x in lines]
    tf = cell if keep_shape else cell.text_frame
    base_p = tf.paragraphs[0]
    proto = copy.deepcopy(base_p._p)
    txBody = tf._txBody

    # 探测模板行前缀（如 '◆ '）：优先取"首个 run 为纯符号"的段落
    prefix = None
    if keep_prefix:
        for para in tf.paragraphs:
            pf = _para_prefix(para._p)
            if pf:
                prefix = pf
                break
        if prefix is None:
            # 回退：直接取第一段首个 run 里的符号字符
            pts = proto.findall(qn("a:r"))
            if pts:
                tn = pts[0].find(qn("a:t"))
                if tn is not None and tn.text:
                    m = re.match(r"^[◆●•►▶·]\s*", tn.text)
                    if m:
                        prefix = m.group(0)

    def set_para_text(p_el, txt):
        ts = p_el.findall(qn("a:r"))
        if ts:
            first = ts[0]
            tnode = first.find(qn("a:t"))
            if tnode is None:
                tnode = first.makeelement(qn("a:t"), {})
                first.append(tnode)
            tnode.text = txt
            for r in ts[1:]:
                p_el.remove(r)
        else:
            pass

    new_ps = []
    for ln in lines:
        pel = copy.deepcopy(proto)
        set_para_text(pel, ln)
        new_ps.append(pel)

    # 为每行补前缀：把前缀作为独立 run 插到段首（保持与模板同构）
    if prefix:
        for pel in new_ps:
            ts = pel.findall(qn("a:r"))
            if not ts:
                continue
            first = ts[0]
            tnode = first.find(qn("a:t"))
            cur = tnode.text if (tnode is not None and tnode.text is not None) else ""
            if cur.strip() == "":
                # 空行不补前缀
                continue
            # 若内容本身已带符号则不重复
            if re.match(r"^[◆●•►▶·]", cur.lstrip()):
                continue
            # 复制首个 run 的样式作为前缀 run
            pref_run = copy.deepcopy(first)
            pt = pref_run.find(qn("a:t"))
            if pt is None:
                pt = pref_run.makeelement(qn("a:t"), {})
                pref_run.append(pt)
            pt.text = prefix
            first.addprevious(pref_run)

    for old in txBody.findall(qn("a:p")):
        txBody.remove(old)
    for pel in new_ps:
        txBody.append(pel)


def set_shape_text(shape, text, lang=None):
    """写普通文本框（保留首个 run 样式）"""
    if not shape.has_text_frame or text is None:
        return
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for r in p0.runs[1:]:
            r.text = ""
    else:
        p0.text = text
    for extra in tf.paragraphs[1:]:
        for r in extra.runs:
            r.text = ""


def get_path(obj, path):
    cur = obj
    for k in path:
        if isinstance(cur, dict):
            cur = cur.get(k)
        else:
            return None
    return cur


def replace_picture(slide, old_pic_shape, img_stream):
    """把已有图片形状换为新图，保持原位置/尺寸。

    ⚠️ 必须用 lxml + QName 直接写 r:embed：
    python-pptx 的 blip.embed setter 对这种模板里的 blip 节点会静默失败
    （赋值不报错但 XML 不变），导致生成的 PPTX 里仍是模板原图。
    """
    try:
        blip = old_pic_shape._element.blipFill.blip
        part = slide.part
        image_part, new_rId = part.get_or_add_image_part(img_stream)
        # 直接操作 XML 属性，绕开 python-pptx 的 setter
        R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        blip.set("{%s}embed" % R_NS, new_rId)
        # 去掉可能存在的外链，避免显示旧图
        for attr in ("{%s}link" % R_NS,):
            if attr in blip.attrib:
                del blip.attrib[attr]
        # 顺便清掉模板的裁切偏移（srcRect），让整图完整显示、不被裁掉一截
        blipFill = old_pic_shape._element.blipFill
        for src in blipFill.findall(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}srcRect"):
            blipFill.remove(src)
        return True
    except Exception as e:
        print("  ! 图片替换失败:", e, file=sys.stderr)
        return False


BADGE_TEXT = {
    "zh": "示意图，以实物为准",
    "en": "For reference only",
}


def add_badge(slide, pic_shape, lang="en"):
    """在图片右下角叠加一个小标签（白底 + 灰字）。

    lang: "zh" 印「示意图，以实物为准」；"en" 印「For reference only」
    """
    text = BADGE_TEXT.get(lang, BADGE_TEXT["en"])
    try:
        from pptx.util import Emu, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

        PW, PH = pic_shape.width, pic_shape.height
        # 英文串更长，宽度按文本长度给
        ratio = 0.44 if lang == "en" else 0.34
        box_w = int(PW * ratio)
        box_h = int(PH * 0.075)
        margin = int(min(PW, PH) * 0.03)
        left = pic_shape.left + PW - box_w - margin
        top = pic_shape.top + PH - box_h - margin

        tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(box_w), Emu(box_h))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = 0; tf.margin_right = 0
        tf.margin_top = 0; tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = text
        run.font.size = Pt(9)
        run.font.bold = False
        run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)   # 灰色文字
        # 白底（略半透明，避免太突兀）
        fill = tb.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        try:
            from pptx.oxml.ns import qn
            srgb = fill.fore_color._xFill.find(qn('a:srgbClr'))
            if srgb is not None:
                alpha = srgb.makeelement(qn('a:alpha'), {'val': '88000'})
                srgb.append(alpha)
        except Exception:
            pass
        tb.line.fill.background()
        return tb
    except Exception as e:
        print("  ! 示意图角标添加失败:", e, file=sys.stderr)
        return None


def fill_pictures(slide, img_map, pic_slots=None, lang="en"):
    """把机组照片填进 slide 的大图位。

    img_map: {"open": 流|None, "silent": 流|None, "badge": {...}}
    pic_slots: [左框shape, 右框shape]（不传则自动取面积最大的两个图位）
    规则：两张都有→左开架右静音；只有一张→居中放大（占两框合并后的中间区域）。
    """
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    if not pics:
        return
    if pic_slots:
        left_pic, right_pic = pic_slots
    else:
        big = sorted(pics, key=lambda s: s.width * s.height, reverse=True)[:2]
        if len(big) < 2:
            return
        # 按 left 排序定左右
        big = sorted(big, key=lambda s: s.left)
        left_pic, right_pic = big[0], big[1]

    def _hide(shape):
        """不删除形状（删了模板里盒子位置会变），只清空图片引用。"""
        try:
            blip = shape._element.blipFill.blip
            blip.getparent().remove(blip)
        except Exception:
            pass

    def _center_one(slide, shape_a, shape_b, stream, badge):
        """只有一张图 -> 隐藏两个原图位，在两框合并区域的中间放一张大图。"""
        _hide(shape_a)
        _hide(shape_b)
        try:
            from pptx.util import Emu
            L = min(shape_a.left, shape_b.left)
            R = max(shape_a.left + shape_a.width, shape_b.left + shape_b.width)
            T = min(shape_a.top, shape_b.top)
            B = max(shape_a.top + shape_a.height, shape_b.top + shape_b.height)
            W, H = R - L, B - T
            # 保持两个原图位的高度，宽度取 4:3（与图库画布一致）
            img_h = H
            img_w = int(img_h * 4 / 3)
            if img_w > W:
                img_w = W
                img_h = int(img_w * 3 / 4)
            left = L + (W - img_w) // 2
            top = T + (H - img_h) // 2
            pic = slide.shapes.add_picture(stream, Emu(left), Emu(top),
                                           Emu(img_w), Emu(img_h))
            if badge:
                add_badge(slide, pic, lang=lang)
            return pic
        except Exception as e:
            print("  ! 单图居中失败:", e, file=sys.stderr)
            return None

    o, s_ = img_map.get("open"), img_map.get("silent")
    badges = img_map.get("badge") or {}

    if o and s_:
        ok_o = replace_picture(slide, left_pic, o)
        ok_s = replace_picture(slide, right_pic, s_)
        if ok_o and badges.get("open"):
            add_badge(slide, left_pic, lang=lang)
        if ok_s and badges.get("silent"):
            add_badge(slide, right_pic, lang=lang)
    elif s_ and not o:
        # 只有静音图 -> 居中
        _center_one(slide, left_pic, right_pic, s_, badges.get("silent"))
    elif o and not s_:
        # 只有开架图 -> 居中
        _center_one(slide, left_pic, right_pic, o, badges.get("open"))


def set_title_on_slide(slide, text, sub_text=None):
    """替换 slide1 顶部主标题 + 小标题（型号字）。

    模板结构（实测）：
      - 小标题（型号字，如 AS500E5/WP715D5）：top ≈ 0.12in，文本不含标题关键词
      - 主标题（大标题）：top ≈ 0.63~0.72in，文本含 'Industrial Range'/'Powered By'/'采用'
    定位规则：先按文本关键词找主标题；小标题取“顶部且非主标题”的那个框。
    """
    MAIN_KEYS = ("Industrial Range", "Powered By", "采用")

    def _is_main(t):
        return any(k in t for k in MAIN_KEYS)

    # 收集顶部（top < 2.0in）的文本框，按 top 升序
    tops = []
    for sh in slide.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        if sh.top is not None and sh.top < 2000000:
            tops.append(sh)
    tops.sort(key=lambda sh: sh.top or 0)

    main_sh = None
    sub_sh = None
    for sh in tops:
        t = sh.text_frame.text.strip()
        if _is_main(t):
            main_sh = sh          # 含关键词 -> 主标题
        elif sub_sh is None:
            sub_sh = sh           # 其余最上面的 -> 小标题（型号字）

    # 兜底：模板关键词变了，就按位置取（最上面=小标题，次上=主标题）
    if main_sh is None and len(tops) >= 2:
        sub_sh, main_sh = tops[0], tops[1]

    if text and main_sh is not None:
        set_shape_text(main_sh, text)
    if sub_text and sub_sh is not None:
        set_shape_text(sub_sh, sub_text)



def _replace_slide_with_bg(slide, img_path, blank_layout_idx=None):
    """2026-10-19 用户规则：第 3、4 页直接换成用户提供的成品整页图片（纯版式，不参与数据填充）。

    做法：清掉该页所有 shapes，插入一张铺满整页的图片（位置/尺寸与用户成品一致）。
    这样导出 PDF 时不会因为往文本框写数据而破坏排版（折行/两端对齐/◆ 前缀）。
    """
    from pptx.util import Emu
    # ⚠️ 必须用页面自身尺寸并从 (0,0) 开始铺满。
    # 之前写死 A4 8.27in 并把 left 也设成 8.27in（错把宽度当左边距），
    # 导致图片被推到页面右边界之外 —— 第 3/4 页看着就是空白。
    prs = slide.part.package.presentation_part.presentation
    W, H = prs.slide_width, prs.slide_height
    sp_tree = slide.shapes._spTree
    for sh in list(slide.shapes):
        sp_tree.remove(sh._element)
    slide.shapes.add_picture(img_path, 0, 0, width=W, height=H)



# 2026-10-19 用户反馈：'@75% load&7m dB(A)' 这类「靠空格右推」的固定标签行，
# 在 Windows/PDF 渲染时字宽略有差异就会把 dB(A) 挤到下一行。
# 做法：出件时把这类框 bodyPr 改成 wrap="none"（不自动换行），文字超出框也不折行。
_NOWRAP_NAMES = ("文本框 19",)


def _set_nowrap(slide):
    from pptx.oxml.ns import qn as _qn
    for sh in slide.shapes:
        if sh.name in _NOWRAP_NAMES and sh.has_text_frame:
            body = sh.text_frame._txBody.find(_qn("a:bodyPr"))
            if body is not None:
                body.set("wrap", "none")


# 2026-10-19 用户要求：燃油消耗量标签（100% 备用(L/H) 等）不许换行。
# 中文模板标签列 1.0236in、减去左右边距只剩 0.8645in，而 8pt 中文实测要 0.779in，
# LibreOffice 渲染又比测量宽，必然折行。做法：出件时把标签列加宽、数值列等量收窄
# （数值只有 3~4 位数字，0.79in 富余很多）。只动列宽，不动行高、字号、文字。
_FUEL_COL_FIX = {
    # 表索引: (标签列号, 目标宽 in, 数值列号, 目标宽 in)
    # 2026-10-20 调整：1.25/0.955 时「Prime & Standby」在 1.18in 的右列里被拆成
    # Standb/y 两行（用户实测反馈）。左右各让 0.15in → 标签列 1.10、数值列 1.105，
    # 实测 8pt 下两列都单行（标签列 8pt 仍放得下）。
    1: (1, 1.10, 2, 1.105),
}


def _fix_fuel_columns(slide, table_tweaks):
    """按 table_map 调整指定表格的列宽，保证燃油标签不折行。"""
    from pptx.util import Inches
    if not table_tweaks:
        return
    tables = [sh for sh in slide.shapes if sh.has_table]
    for idx, (lab_c, lab_w, val_c, val_w) in table_tweaks.items():
        if idx >= len(tables):
            continue
        tbl = tables[idx].table
        try:
            tbl.columns[lab_c].width = Inches(lab_w)
            tbl.columns[val_c].width = Inches(val_w)
        except (IndexError, KeyError):
            pass


# 2026-10-19 用户反馈「导出卡住」：新中文模板把 11 个字体嵌进 PPT（28.4MB），
# 生成件 PPTX 膨胀到 53MB，前端下载迟迟不结束。模板已内嵌字体是给编辑用的，
# 出件不需要 —— 剥掉 ppt/fonts/* + presentation.xml 里的 embeddedFontLst。
def _strip_embedded_fonts(pptx_path):
    """删除 PPTX 内嵌字体，显著减小文件（不影响 PDF 转换与显示）。"""
    import zipfile as _zf, shutil as _sh, tempfile as _tf, re as _re, os as _os
    try:
        zin = _zf.ZipFile(pptx_path)
    except Exception:
        return
    tmp = pptx_path + ".nofont.tmp"
    try:
        zout = _zf.ZipFile(tmp, "w", _zf.ZIP_DEFLATED)
        for item in zin.infolist():
            if item.filename.startswith("ppt/fonts/"):
                continue
            data = zin.read(item.filename)
            if item.filename in ("ppt/presentation.xml", "[Content_Types].xml"):
                txt = data.decode("utf-8", "ignore")
                txt = _re.sub(r"<p:embeddedFontLst>.*?</p:embeddedFontLst>", "", txt, flags=_re.S)
                txt = _re.sub(r"<Override[^>]*ppt/fonts/[^>]*/>", "", txt)
                txt = _re.sub(r'<Default Extension="fntdata"[^>]*/>', "", txt)
                data = txt.encode("utf-8")
            zout.writestr(item, data)
        zout.close()
        zin.close()
        _sh.move(tmp, pptx_path)
    except Exception as e:
        try:
            zin.close()
        except Exception:
            pass
        if _os.path.exists(tmp):
            _os.remove(tmp)
        print("  警告: 剥离内嵌字体失败(%s)，保留原文件" % e)


# 2026-10-19 用户反馈「导出还卡」：模板第3/4页背景图 + 机组照片全是原始大图，
# 单张 11MB。出件前统一压到长边 1600px / JPEG q88（A4 打印 300dpi 绰绰有余），
# 文件从 22.8MB 降到几 MB，下载和转换都快。
def _compress_media(pptx_path, max_side=1600, quality=88):
    import zipfile as _zf, shutil as _sh, os as _os, io as _io
    try:
        from PIL import Image
    except Exception:
        return
    try:
        zin = _zf.ZipFile(pptx_path)
    except Exception:
        return
    tmp = pptx_path + ".small.tmp"
    changed = 0
    try:
        zout = _zf.ZipFile(tmp, "w", _zf.ZIP_DEFLATED)
        for item in zin.infolist():
            data = zin.read(item.filename)
            low = item.filename.lower()
            if low.startswith("ppt/media/") and low.endswith((".png", ".jpg", ".jpeg")):
                try:
                    im = Image.open(_io.BytesIO(data))
                    im.load()
                    w, h = im.size
                    if max(w, h) > max_side or len(data) > 600 * 1024:
                        if im.mode in ("RGBA", "LA", "P"):
                            bg = Image.new("RGB", im.size, (255, 255, 255))
                            conv = im.convert("RGBA")
                            bg.paste(conv, mask=conv.split()[-1])
                            im = bg
                        else:
                            im = im.convert("RGB")
                        im.thumbnail((max_side, max_side), Image.LANCZOS)
                        buf = _io.BytesIO()
                        im.save(buf, "JPEG", quality=quality, optimize=True)
                        new = buf.getvalue()
                        if len(new) < len(data):
                            data = new
                            changed += 1
                except Exception:
                    pass
            zout.writestr(item, data)
        zout.close()
        zin.close()
        _sh.move(tmp, pptx_path)
        if changed:
            print("  压缩媒体 %d 张" % changed)
    except Exception as e:
        try:
            zin.close()
        except Exception:
            pass
        if _os.path.exists(tmp):
            _os.remove(tmp)
        print("  警告: 媒体压缩失败(%s)" % e)


# ===== 出件前单位归一（2026-10-20 用户定：后端也补一份，双保险）=====
# 前端 index.html normalizeUnits() 的 Python 镜像。
# 规则：值末尾已带单位就不重复补；值是 TBA 一律不补单位。
# 前端漏补 / 直接调 API 时也不会丢单位。
_UNIT_SUFFIX = [
    # (字段, 单位, 是否空格, 已带单位判定)
    ("prp_kva", "kVA", True, r"kVA$"),
    ("eps_kva", "kVA", True, r"kVA$"),
    ("prp_kw", "kW", True, r"kW$"),
    ("eps_kw", "kW", True, r"kW$"),
    ("o_wt", "kg", True, r"kg$"),
    ("s_wt", "kg", True, r"kg$"),
    ("o_tank", "L", True, r"L$"),
    ("s_tank", "L", True, r"L$"),
    ("displ", "L", True, r"L$"),
    ("rpm", "rpm", True, r"rpm$"),
    ("ex_t", "°C", True, r"(°C|℃)$"),
    ("in_t", "°C", True, r"(°C|℃)$"),
    ("co_oil", "L", True, r"L$"),
    ("co_cool", "L", True, r"L$"),
]

# 中文下拉项 -> 英文源值（与前端 OPT_ZH 一一对应；后端按 lang 自行反查）
_OPT_ZH = {
    "Self-excited": "自励磁", "PMG": "永磁(PMG)", "Separately excited": "他励",
    "Star": "星形(Y)", "Delta": "三角形(Δ)",
    "Two thirds": "2/3", "Five sixths": "5/6",
    "Single phase": "单相", "Three phase": "三相",
    "Single bearing": "单轴承", "Double bearing": "双轴承",
}
_ZH2EN = {v: k for k, v in _OPT_ZH.items()}


def normalize_units(src):
    """把「数字 + 固定单位」补成模板要的完整字符串（幂等，重复导出不会叠加）。"""
    d = dict(src or {})
    for k, unit, space, already in _UNIT_SUFFIX:
        v = str(d.get(k) if d.get(k) is not None else "").strip()
        if not v:
            continue
        if re.sub(r"\s", "", v).upper() == "TBA":
            continue                      # TBA 不补单位
        if re.search(already, v, re.I):
            continue                      # 已带单位不重复补
        d[k] = v + (" " if space else "") + unit
    # 起动机电压：24 -> DC24V
    v = str(d.get("el_v") or "").strip().upper().replace(" ", "")
    if v and v != "TBA":
        if not v.startswith("DC"):
            v = "DC" + v
        if not v.endswith("V"):
            v += "V"
        d["el_v"] = v
    # 绝缘/温升：Class H/H（绝缘固定 H，只取温升段）
    g = str(d.get("g_insul") or "").strip()
    if g and g.upper() != "TBA":
        mm = re.match(r"^\s*Class\s*[A-Z]\s*/\s*([A-Z])\s*$", g, re.I)
        if mm:
            rise = mm.group(1).upper()
        elif re.fullmatch(r"[A-Z]", g, re.I):
            rise = g.upper()
        else:
            rise = g
        d["g_insul"] = "Class H/" + rise
    return d


_OPT_KEYS = ("g_exc", "g_conn", "g_pitch", "g_phase", "g_brg")


def localize_options(d, lang):
    """下拉项按语言转换（前端 optLabel 的镜像）。"""
    for k in _OPT_KEYS:
        v = d.get(k)
        if not v:
            continue
        if lang == "zh" and v in _OPT_ZH:
            d[k] = _OPT_ZH[v]
        elif lang == "en" and v in _ZH2EN:
            d[k] = _ZH2EN[v]
    return d


def build(data, lang="en", template=None, outdir="output", images_dir=None, transparent_bg=False, pic_slots=None):
    """生成参数单 PPTX。

    2026-10-20 用户定「双保险」：进 build 先做一遍单位归一 + 下拉翻译，
    前端已补过也不会重复补（幂等），直接调 API 也有完整单位。
    """
    data = localize_options(normalize_units(data), lang)
    tmap = load_map(lang)
    tpl = template or os.path.join(HERE, tmap["template"])
    prs = Presentation(tpl)
    images_dir = images_dir or os.path.join(HERE, "images")

    for slide_def in tmap["slides"]:
        s_idx = slide_def["slide"] - 1
        slide = prs.slides[s_idx]
        # 1 页：声压级标签框改不换行（防 dB(A) 掉行）
        if s_idx == 0:
            _set_nowrap(slide)
            # 1 页：燃油消耗量标签列加宽，防「100% 备用(L/H)」折行
            _fix_fuel_columns(slide, _FUEL_COL_FIX)

        # 第 3/4 页：直接替换为用户成品整页图（纯版式，不填数据）
        if slide_def.get("background"):
            bg = os.path.join(HERE, slide_def["background"])
            if os.path.exists(bg):
                _replace_slide_with_bg(slide, bg)
                print("  slide%d -> 整页背景图 %s" % (s_idx + 1, slide_def["background"]))
                continue
        # 机组照片（仅 slide1）
        if s_idx == 0 and images_dir:
            import image_utils
            found = image_utils.find_images(images_dir, data.get("model", ""))
            # 手动指定图优先（data.__img_open__ / __img_silent__ 为绝对路径）
            at = data.get("__img_open__")
            st = data.get("__img_silent__")
            if at:
                found["open"] = at
            if st:
                found["silent"] = st
            if found.get("open") or found.get("silent"):
                img_map = {}
                # 先用占位框尺寸算裁切比（左框）
                pics = sorted([sh for sh in slide.shapes if sh.shape_type == 13],
                              key=lambda s: s.width * s.height, reverse=True)[:2]
                if len(pics) >= 2:
                    bw, bh = pics[0].width, pics[0].height
                else:
                    bw, bh = 3000000, 2000000
                # 示意图标记：前端传的 __img_badge__ + 后端兜底套用的自动标记
                badge_req = dict(found.get("_badge") or {})
                front_badge = data.get("__img_badge__") or {}
                if isinstance(front_badge, dict):
                    for bk, bv in front_badge.items():
                        if bv:
                            badge_req[bk] = True
                badges = {}
                for key in ("open", "silent"):
                    p = found.get(key)
                    if p and os.path.exists(p):
                        img_map[key] = image_utils.prepare_image(
                            p, bw, bh, transparent_bg=transparent_bg)
                        if badge_req.get(key):
                            badges[key] = True
                if badges:
                    img_map["badge"] = badges
                if img_map.get("open") or img_map.get("silent"):
                    fill_pictures(slide, img_map, pic_slots, lang=lang)
                    print("  已插入图片:", [k for k in img_map if k != "badge"],
                          ("(示意图: %s)" % ",".join(badges)) if badges else "")

        # 标题：仅第 1 页替换（第 2/3/4 页保持模板原样）
        title_tpl = tmap.get("title") if s_idx == 0 else None
        sub_title = data.get("__title_sub__") if s_idx == 0 else None
        title_custom = ((data.get("__title_custom__") or "").strip()) if s_idx == 0 else ""
        if title_tpl or sub_title or title_custom:
            brand_raw = data.get("e_brand") or ""
            brand = translate(brand_raw, lang) if brand_raw else ""
            if title_custom:
                # 用户手填大标题优先（支持 {engine_brand} 占位）
                title_text = title_custom.replace("{engine_brand}", str(brand))
            elif title_tpl:
                title_text = title_tpl.replace("{engine_brand}", str(brand))
            else:
                title_text = None
            set_title_on_slide(slide, title_text, sub_title)
        # 独立文本框字段（如 slide1 右下角频率/燃料小格、声压级符号框）
        for sdef in slide_def.get("shapes", []):
            si = sdef.get("shape")
            if si is None or si >= len(slide.shapes):
                print("  ! slide%d shape idx=%s 不存在，跳过" % (s_idx + 1, si), file=sys.stderr)
                continue
            sh = slide.shapes[si]
            if not sh.has_text_frame:
                continue
            fname = sdef.get("field")
            val = data.get(fname)
            if val in (None, ""):
                continue
            if sdef.get("mode") == "list":
                lines = val if isinstance(val, (list, tuple)) else [x for x in str(val).split("\n")]
                set_multiline(sh.text_frame, lines, keep_shape=True)
            else:
                set_shape_text(sh, translate(val, lang))

        shape_tables = [sh for sh in slide.shapes if sh.has_table]
        for t_def in slide_def.get("tables", []):
            if t_def["idx"] >= len(shape_tables):
                print("  ! slide%d table idx=%d 不存在，跳过" % (s_idx + 1, t_def["idx"]), file=sys.stderr)
                continue
            table = shape_tables[t_def["idx"]].table
            # 普通字段
            for fname, coord in t_def.get("fields", {}).items():
                val = get_path(data, fname.split(".")) if "." in fname else data.get(fname)
                if val in (None, ""):
                    continue
                r, c = coord
                size_pt = (t_def.get("cell_size") or {}).get(fname)
                txt = translate(val, lang)
                # 单元格级强制译法（中文件用中文、英文件用英文，见 template_map.*.json 的 cell_value_map）
                vmap = (t_def.get("cell_value_map") or {}).get(fname)
                if vmap and str(txt) in vmap:
                    txt = vmap[str(txt)]
                # 带单位字段：数字后自动补单位（如重量 kg / 油箱 L）
                usfx = (t_def.get("cell_suffix") or {}).get(fname)
                if usfx and str(txt).strip() and not re.search(r"[A-Za-z℃°]$", str(txt).strip()):
                    txt = "%s %s" % (str(txt).strip(), usfx)
                # 指定字段禁止换行（长尺寸串 / 带 kVA、kW 的功率值不要拆成两行）
                nw = fname in (t_def.get("nowrap_fields") or [])
                set_cell(table.cell(r, c), txt, size_pt=size_pt, nowrap=nw)
            # 🚫 单元格一律不许折行（2026-10-20 用户要求：中文件用中文、英文件用英文，都不换行）
            # 收尾统一过一遍：全部单元格按可用宽度缩字号，窄列优先整宽单行，
            # 实在放不下再按「单词」折（CJK 无空格，自然整段不拆）。
            if t_def.get("purpose") in NOWRAP_PURPOSES:
                _nowrap_table(table)
            # 【2026-10-20 用户定】没数据的格子一律写 TBA（不带单位），
            # 绝不让模板占位 12345 露到成品上（TBA 兜底）。
            # 只动「模板原文/剩余内容仍是纯数字占位」的格子，表头/标签列一字不碰。
            # 油耗格（多行）由下面的 fuel_consumption 分支处理，这里跳过。
            fc_cell = (t_def.get("fuel_consumption") or {}).get("cell")
            # 【2026-10-23 修复】不能只按「纯数字」判定占位：
            # 用户填的真实数据也可能是纯数字（如绕组出线 12、出线数 6），
            # 会被误判成模板占位 12345 → 写成 TBA。
            # 正确判定：①该格是本表某个字段的写入位置 且 ②data 里该字段有值
            #          → 是真实数据，绝不覆盖。
            # 只有「模板留下的占位且没有对应数据写入」才兜底 TBA。
            _field_cells = {}
            for _fname, _coord in (t_def.get("fields") or {}).items():
                _v = get_path(data, _fname.split(".")) if "." in _fname else data.get(_fname)
                _field_cells[(_coord[0], _coord[1])] = (_v not in (None, ""))
            for r in range(len(table.rows)):
                for c in range(len(table.columns)):
                    if fc_cell and [r, c] == list(fc_cell):
                        continue
                    cell = table.cell(r, c)
                    if not re.fullmatch(r"\s*\d+\s*", cell.text or ""):
                        continue
                    # 这一格是字段写入位且该字段有真实数据 → 保留，不写 TBA
                    if _field_cells.get((r, c)):
                        continue
                    set_cell(cell, "TBA")
            # 多行字段（油耗）
            if "fuel_consumption" in t_def:
                fc = t_def["fuel_consumption"]
                lines = data.get(fc["data_key"])
                if isinstance(lines, str):
                    lines = [x for x in lines.split("\n") if x.strip()]
                if lines and any(x not in (None, "") for x in lines):
                    r, c = fc["cell"]
                    set_multiline(table.cell(r, c), lines)
            # 多行文本字段（features）
            for fname, coord in t_def.get("multiline_fields", {}).items():
                raw = data.get(fname)
                if not raw:
                    continue
                lines = [x for x in str(raw).split("\n") if x.strip()]
                if lines:
                    r, c = coord
                    set_multiline(table.cell(r, c), [translate(x, lang) for x in lines],
                                  keep_prefix=True)

    os.makedirs(os.path.join(HERE, outdir) if not os.path.isabs(outdir) else outdir, exist_ok=True)
    model = data.get("model", "model")
    outpath = os.path.join(os.path.join(HERE, outdir) if not os.path.isabs(outdir) else outdir,
                           "%s_paramsheet_%s.pptx" % (model, lang))
    prs.save(outpath)
    # 剥掉模板内嵌字体（约 28MB），加速前端下载与服务器转换
    _strip_embedded_fonts(outpath)
    # 压缩媒体大图（单张 11MB -> 几百 KB）
    _compress_media(outpath)
    return outpath


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", required=True)
    ap.add_argument("--lang", default="en")
    ap.add_argument("--template")
    ap.add_argument("--out", default="output")
    a = ap.parse_args()
    with io.open(a.data, encoding="utf-8") as f:
        data = json.load(f)
    out = build(data, a.lang, a.template, a.out)
    print(out)


if __name__ == "__main__":
    main()
