#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
批量从源参数单 PPTX 第 1 页抽出机组图片
------------------------------------------------
用法（公司电脑）：
    python extract_images.py --src "D:\\AGG-机组参数单" --out "C:\\Users\\Administrator\\Desktop\\机型图片"

逻辑：
  * 只看第 1 页（机组图片就在第 1 页）
  * 按面积排序取最大的 2 张图
  * 左边那张 = 开架 OPEN，右边那张 = 静音 SILENT
  * 用文件名左侧的机组型号命名：AS2500E5_OPEN.png / AS2500E5_SILENT.png
  * 文件名带「旧版本」的跳过
  * 同时输出 index.json 记录：每份文件抽到几张图、尺寸、是否疑似模板占位图

产出：
  <out>/OPEN/<机型>_OPEN.png
  <out>/SILENT/<机型>_SILENT.png
  <out>/index.json
  <out>/report.txt
"""
import argparse
import hashlib
import json
import os
import re
import sys
import time

# ---------- 机组型号提取（与 convert_extract.py 保持一致） ----------
_UNIT_MODEL_RE = re.compile(r"^[A-Z]{1,3}\d{1,5}[DEG]?[56]/?[A-Z0-9]{0,4}$", re.I)
_UNIT_BAD_WORDS = re.compile(
    r"模板|模版|样板|照片|图片|英文版|中文版|俄语版|旧|Engine|HP$|TBA|^VPS$", re.I)
_OLD_RE = re.compile(r"旧\s*版本|旧版|旧\s*DG|^旧\s|\s-旧|旧图|旧参数", re.I)


def split_unit_engine(src):
    """文件名 -> (机组型号, 发动机型号)；拆不出返回 ('','')"""
    fn = os.path.basename((src or "").replace("\\", "/"))
    fn = re.sub(r"\.(pptx|ppt|pdf)$", "", fn, flags=re.I)
    fn = re.sub(r"[(（][^)）]*[)）]", " ", fn)
    fn = re.sub(r"^(旧|新)\s*", "", fn.strip())
    fn = re.sub(r"^DG\s*SPEC\s+", "", fn, flags=re.I).strip()
    first = fn.split()[0] if fn.split() else ""
    if "-" not in first:
        return "", ""
    segs = first.split("-")
    phase = ""
    if len(segs) > 2 and re.fullmatch(r"[13]P", segs[1].strip(), re.I):
        phase = segs[1].strip().upper()
        segs = [segs[0]] + segs[2:]
    first = "-".join(segs)
    left, right = first.split("-", 1)
    left, right = left.strip(), right.strip()
    if right and " " in right:
        right = right.split()[0]
    if not left or not right:
        return "", ""
    if _UNIT_BAD_WORDS.search(left) or not _UNIT_MODEL_RE.match(left):
        return "", ""
    if phase and "/" not in left:
        left = left + "/" + phase
    return left.upper(), right.upper()


_SAFE_RE = re.compile(r'[\\/:*?"<>|]')


def safe_name(name):
    """型号转安全文件名：/ \\ : * ? " < > | 全部替换成 -（单相 /1P 会变成 -1P）"""
    s = _SAFE_RE.sub("-", str(name or "").strip())
    s = re.sub(r"-{2,}", "-", s).strip("-. ")
    return s or "未命名"


def is_old_version(src):
    fn = os.path.basename((src or "").replace("\\", "/"))
    return bool(_OLD_RE.search(fn))


# ---------- 抽图 ----------
def extract_from_pptx(path, min_side_in=1.2, max_pics=2):
    """第 1 页 -> [(image, left, top, w, h)]，按面积从大到小，最多 max_pics 张"""
    from pptx import Presentation
    prs = Presentation(path)
    if not prs.slides:
        return []
    s = prs.slides[0]
    pics = []
    for sh in s.shapes:
        if sh.shape_type != 13:          # 13 = PICTURE
            continue
        try:
            im = sh.image
        except Exception:
            continue
        w_in = (sh.width or 0) / 914400.0
        h_in = (sh.height or 0) / 914400.0
        # 过滤小图标 / logo：最短边小于 min_side_in 的丢掉
        if min(w_in, h_in) < min_side_in:
            continue
        pics.append({
            "blob": im.blob, "ext": (im.ext or "png").lower(),
            "left": (sh.left or 0) / 914400.0, "top": (sh.top or 0) / 914400.0,
            "w": w_in, "h": h_in, "area": w_in * h_in,
        })
    pics.sort(key=lambda x: -x["area"])
    return pics[:max_pics]


def collect_pptx(src):
    out = []
    for root, _d, files in os.walk(src):
        for fn in files:
            if fn.lower().endswith((".pptx", ".ppt")):
                out.append(os.path.join(root, fn))
    out.sort()
    return out


def main():
    ap = argparse.ArgumentParser()
    # 默认值：不传参也能直接跑（F5 调试 / 双击 bat / python extract_images.py）
    _here = os.path.dirname(os.path.abspath(__file__))
    _def_src = r"D:\AGG-机组参数单"
    if not os.path.isdir(_def_src):
        for cand in (r"D:\AGG", _here, os.path.join(_here, "..")):
            if os.path.isdir(os.path.join(cand, "AGG-机组参数单")):
                _def_src = os.path.join(cand, "AGG-机组参数单"); break
    ap.add_argument("--src", default=_def_src, help="源目录，默认 D:\\AGG-机组参数单")
    ap.add_argument("--out", default=os.path.join(_here, "机型图片"), help="输出目录，默认脚本同目录下「机型图片」")
    ap.add_argument("--min-side", type=float, default=1.2,
                    help="图片最短边下限（英寸），小于此值当图标忽略，默认 1.2")
    ap.add_argument("--limit", type=int, default=0, help="只处理前 N 个（调试用）")
    ap.add_argument("--no-recursive", action="store_true")
    a = ap.parse_args()

    src, out = a.src, a.out
    if not os.path.isdir(src):
        print("源目录不存在:", src, file=sys.stderr)
        return 2

    os.makedirs(os.path.join(out, "OPEN"), exist_ok=True)
    os.makedirs(os.path.join(out, "SILENT"), exist_ok=True)

    files = collect_pptx(src) if not a.no_recursive else [
        os.path.join(src, f) for f in sorted(os.listdir(src))
        if f.lower().endswith((".pptx", ".ppt"))]
    if a.limit:
        files = files[:a.limit]

    print("共 %d 份 PPTX 待处理" % len(files))
    t0 = time.time()
    index, rows = [], []
    n_ok = n_old = n_nopics = n_nomodel = n_err = 0
    # 用于识别"模板占位图"：同样的图片内容出现在多个机型里 -> 多半是模板图
    hash_count = {}

    for i, p in enumerate(files, 1):
        rel = os.path.relpath(p, src)
        if is_old_version(rel):
            n_old += 1
            continue
        unit, engine = split_unit_engine(rel)
        if not unit:
            n_nomodel += 1
            rows.append({"src": rel, "status": "no_model"})
            continue
        try:
            pics = extract_from_pptx(p, a.min_side)
        except Exception as e:
            n_err += 1
            rows.append({"src": rel, "status": "error", "error": str(e)})
            continue
        if not pics:
            n_nopics += 1
            rows.append({"src": rel, "model": unit, "status": "no_picture"})
            continue

        # 左边=开架，右边=静音；只有一张就按位置判断（偏左=开架）
        pics.sort(key=lambda x: x["left"])
        saved = []
        for k, pc in enumerate(pics):
            if len(pics) == 1:
                kind = "OPEN" if pc["left"] < 3.0 else "SILENT"
            else:
                kind = "OPEN" if k == 0 else "SILENT"
            h = hashlib.md5(pc["blob"]).hexdigest()
            hash_count[h] = hash_count.get(h, 0) + 1
            safe_unit = safe_name(unit)
            fn = "%s_%s.%s" % (safe_unit, kind, pc["ext"])
            dst = os.path.join(out, kind, fn)
            n = 2
            base = os.path.join(out, kind, "%s_%s" % (safe_unit, kind))
            while os.path.exists(dst):
                dst = "%s_%d.%s" % (base, n, pc["ext"])
                n += 1
            try:
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                with open(dst, "wb") as f:
                    f.write(pc["blob"])
            except Exception as e:
                rows.append({"src": rel, "model": unit, "status": "write_error", "error": str(e)})
                continue
            saved.append({"kind": kind, "file": os.path.relpath(dst, out),
                          "w": round(pc["w"], 2), "h": round(pc["h"], 2),
                          "kb": len(pc["blob"]) // 1024, "hash": h})
        n_ok += 1
        index.append({"model": unit, "engine": engine, "src": rel, "images": saved})
        rows.append({"src": rel, "model": unit, "status": "ok", "n": len(saved)})
        if i % 100 == 0:
            print("  ...已处理 %d/%d" % (i, len(files)))

    # 标出疑似模板占位图（同一张图出现在 >3 个机型里）
    dup = {h for h, c in hash_count.items() if c > 3}
    for it in index:
        for im in it["images"]:
            im["suspect_template"] = im["hash"] in dup

    # ---------- 按内容去重：同一张图只在 DEDUP/ 里存一份 ----------
    os.makedirs(os.path.join(out, "DEDUP"), exist_ok=True)
    hash_file = {}       # hash -> dedup 里的文件名
    hash_models = {}     # hash -> [机型...]
    for it in index:
        for im in it["images"]:
            h = im["hash"]
            hash_models.setdefault(h, []).append("%s(%s)" % (it["model"], im["kind"]))
            if h in hash_file:
                im["dedup_file"] = hash_file[h]
                continue
            dst = os.path.join(out, "DEDUP", im["file"].replace(os.sep, "_"))
            src_p = os.path.join(out, im["file"])
            try:
                if os.path.exists(src_p):
                    import shutil
                    shutil.copy2(src_p, dst)
                    hash_file[h] = "DEDUP/" + os.path.basename(dst)
                    im["dedup_file"] = hash_file[h]
            except Exception:
                pass

    # 去重清单：只含「一张图 -> 用它的机型」，便于挑选
    manifest = [{"dedup_file": hash_file[h], "count": len(v), "models": v}
                for h, v in hash_models.items() if h in hash_file]
    manifest.sort(key=lambda x: -x["count"])
    with open(os.path.join(out, "dedup_manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)

    with open(os.path.join(out, "index.json"), "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=1)

    secs = time.time() - t0
    with open(os.path.join(out, "report.txt"), "w", encoding="utf-8") as f:
        f.write("参数单图片抽取报告\n")
        f.write("=" * 50 + "\n")
        f.write("源目录: %s\n输出目录: %s\n" % (src, out))
        f.write("耗时: %.1f 秒\n\n" % secs)
        f.write("文件总数    : %d\n" % len(files))
        f.write("成功抽到图  : %d\n" % n_ok)
        f.write("跳过(旧版本): %d\n" % n_old)
        f.write("无图片      : %d\n" % n_nopics)
        f.write("型号识别不出: %d\n" % n_nomodel)
        f.write("读取失败    : %d\n" % n_err)
        f.write("\n去重后唯一图片: %d 张（DEDUP 目录，可直接打包这个目录）\n" % len(hash_file))
        f.write("疑似模板占位图（同一张图出现在多个机型）: %d 张\n" % len(dup))
        if dup:
            f.write("  ⚠ 这些图很可能是模板自带占位图，不是真机组照片\n")
            for h in list(dup)[:10]:
                f.write("    hash=%s  出现 %d 次\n" % (h[:12], hash_count[h]))
        f.write("\n" + "=" * 50 + "\n明细:\n")
        for r in rows:
            f.write("  [%s] %s %s\n" % (r["status"], r.get("model", ""), r["src"]))

    print("\n完成，耗时 %.1f 秒" % secs)
    print("成功抽到图 %d 份 / 无图片 %d / 型号识别不出 %d / 旧版本跳过 %d / 失败 %d"
          % (n_ok, n_nopics, n_nomodel, n_old, n_err))
    print("去重后唯一图片: %d 张 -> %s" % (len(hash_file), os.path.join(out, "DEDUP")))
    print("疑似模板占位图: %d 张（同一张图出现在多个机型）" % len(dup))
    print("输出目录:", out)
    print("\n下一步：把整个输出目录打包成 zip 发回来。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
